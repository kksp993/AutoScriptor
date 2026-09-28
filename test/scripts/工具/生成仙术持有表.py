from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_SKILLS_PATH = PROJECT_ROOT / "skills.txt"
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "仙术持有情况.xlsx"

# 内置完整布局骨架：系列 → (属性, 主动技能, 被动技能)。
# 生成时不再读取任何 Excel 模板，唯一外部数据源是 skills.txt。
SERIES_META = (
    (
        "龙年仙术",
        (
            ("金", "乾坤一掷", "疾速之刃"),
            ("木", "七星连矢", "毒木异变"),
            ("水", "水华漫天", "刺骨之寒"),
            ("火", "旋转火环", "爆燃之火"),
            ("土", "砂暴陷阱", "流沙覆甲"),
            ("雷", "九霄雷劫", "无极静电"),
            ("月", "月曜漫射", "暗月蚀影"),
            ("时", "时空错乱", "凝滞之痕"),
        ),
    ),
    (
        "蛇年仙术",
        (
            ("金", "锋金护身", "锋利反噬"),
            ("木", "荆棘穿刺", "棘刺蔓生"),
            ("水", "怒涛猛进", "叠浪狂潮"),
            ("火", "火焰宝珠", "暗炎之殇"),
            ("土", "岩拳裂地", "地灵守护"),
            ("雷", "雷蛇闪动", "麻痹之触"),
            ("月", "月华普照", "月影庇护"),
            ("时", "时序扭曲", "时空印记"),
        ),
    ),
    (
        "小小英雄",
        (
            ("金", "剑碎尘心", "暗金狩猎"),
            ("木", "圣树送葬", "即刻腐化"),
            ("水", "覆海龙泉", "霜寒之躯"),
            ("火", "灭尽残炎", "怒焰神功"),
            ("土", "地动咆哮", "镇岳元诀"),
            ("雷", "五雷天箓", "霹雳雷诀"),
            ("月", "月神泪", "九阴神诀"),
            ("时", "时空幻剑", "时空紊乱"),
        ),
    ),
    (
        "马年仙术",
        (
            ("金", "金阙断狱", "金刀铸形"),
            ("木", "竹影林侍", "剧毒暴击"),
            ("水", "镇海龙潮", "霜华凝萃"),
            ("火", "烈焰风轮", "狂焰追击"),
            ("土", "崩土连岳", "攻防互易"),
            ("雷", "天谴奔雷", "雷心跃动"),
            ("月", "冥月幽刃", "广寒星陨"),
            ("时", "时界归影", "时缚回响"),
        ),
    ),
    (
        "须弥仙术",
        (
            ("金", "金磁元刹", "金磁破灵"),
            ("木", "树神花禁语", "树神祝福"),
            ("水", "重水乾坤", "重水业障"),
            ("火", "炎焱火莲", "火生莲"),
            ("土", "乾震地脉", "震荡操术"),
            ("雷", "雷渊呼唤", "雷爆"),
            ("月", "月漩光", "月漩之心"),
            ("时", "时空灵固", "时空灵噬"),
        ),
    ),
)

SERIES_ORDER = tuple(series_name for series_name, _ in SERIES_META)
ATTR_ORDER = ("金", "木", "水", "火", "土", "雷", "月", "时")
GRADE_NAMES = ("下", "中", "上", "绝")

# 模板里的旧写法 → skills.txt 实际写法（OCR 或翻译变体）。
TEMPLATE_BASE_NAME_ALIASES = {
    "灭尽残炎": "灭烬残炎",
    "竹影林侍": "竹影灵侍",
    "雷爆": "雷暴",
}

# 布局：A 系列 / B 属性 / C 主动 / D-G 下中上绝 / H 备注 / I 属性 / J 被动 / K-N 下中上绝 / O 备注
ACTIVE_NAME_COLUMN = 3
ACTIVE_GRADE_FIRST_COLUMN = 4
PASSIVE_ATTR_COLUMN = 9
PASSIVE_NAME_COLUMN = 10
PASSIVE_GRADE_FIRST_COLUMN = 11
NOTE_COLUMNS = (8, 15)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")
SERIES_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
OWNED_CHAIN_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
EMPTY_FILL = PatternFill(fill_type=None)
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
LEFT_ALIGNMENT = Alignment(horizontal="left", vertical="center")


def normalize_skill_name(value: str) -> str:
    """归一化 OCR 标点与已知变体字，用于匹配。"""
    return (
        "".join(value.split())
        .replace("・", "·")
        .replace("•", "·")
        .replace("疊", "叠")
    )


def load_owned_grades(skills_path: Path) -> dict[str, set[str]]:
    """把 skills.txt（形如「七星连矢·上」）解析为基础名 → 已拥有等级集合。"""
    owned_grades: dict[str, set[str]] = {}
    for line in skills_path.read_text(encoding="utf-8").splitlines():
        normalized = normalize_skill_name(line)
        if not normalized:
            continue
        if "·" in normalized:
            base_name, grade_name = normalized.rsplit("·", 1)
        else:
            base_name, grade_name = normalized, "已习得"
        owned_grades.setdefault(base_name, set()).add(grade_name)
    return owned_grades


def canonical_base_name(template_base_name: str) -> str:
    """模板旧名 → skills.txt 实际名。"""
    return TEMPLATE_BASE_NAME_ALIASES.get(template_base_name, template_base_name)


def _grade_ownership(
    owned_grades: dict[str, set[str]],
    template_base_name: str,
) -> list[bool]:
    actual_grades = owned_grades.get(canonical_base_name(template_base_name), set())
    return [grade_name in actual_grades for grade_name in GRADE_NAMES]


def _write_skill_block(
    worksheet,
    row_number: int,
    start_name_column: int,
    start_grade_column: int,
    attr: str,
    template_base_name: str,
    owned_grades: dict[str, set[str]],
    stats: list[int],
) -> None:
    worksheet.cell(row_number, start_name_column - 1).value = attr
    worksheet.cell(row_number, start_name_column).value = template_base_name
    ownership = _grade_ownership(owned_grades, template_base_name)
    for grade_offset, is_owned in enumerate(ownership):
        cell = worksheet.cell(row_number, start_grade_column + grade_offset)
        if not is_owned:
            continue
        cell.value = "✔"
        stats[0] += 1
    for grade_offset, is_owned in enumerate(ownership):
        if not is_owned:
            break
        worksheet.cell(
            row_number,
            start_grade_column + grade_offset,
        ).fill = OWNED_CHAIN_FILL
        stats[1] += 1


def _known_series_keys() -> set[str]:
    keys = set()
    for _, rows in SERIES_META:
        for attr, active_name, passive_name in rows:
            keys.add(canonical_base_name(active_name))
            keys.add(canonical_base_name(passive_name))
    return keys


def _style_header(worksheet) -> None:
    headers = (
        ("A1", "系列"), ("B1", "属性"), ("C1", "主动"),
        ("D1", "下"), ("E1", "中"), ("F1", "上"), ("G1", "绝"),
        ("H1", None),
        ("I1", "属性"), ("J1", "被动"),
        ("K1", "下"), ("L1", "中"), ("M1", "上"), ("N1", "绝"),
        ("O1", None),
    )
    for coordinate, _ in headers:
        cell = worksheet[coordinate]
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = CENTER_ALIGNMENT
    for column_number in range(1, 16):
        worksheet.column_dimensions[get_column_letter(column_number)].width = 8
    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 6
    worksheet.column_dimensions["C"].width = 14
    worksheet.column_dimensions["I"].width = 6
    worksheet.column_dimensions["J"].width = 14
    worksheet.freeze_panes = "A2"


def _write_series_section(
    worksheet,
    start_row: int,
    series_name: str,
    series_rows,
    owned_grades: dict[str, set[str]],
    stats: list[int],
) -> int:
    row_number = start_row
    section_start_row = start_row
    for attr, active_name, passive_name in series_rows:
        _write_skill_block(
            worksheet,
            row_number,
            ACTIVE_NAME_COLUMN,
            ACTIVE_GRADE_FIRST_COLUMN,
            attr,
            active_name,
            owned_grades,
            stats,
        )
        _write_skill_block(
            worksheet,
            row_number,
            PASSIVE_NAME_COLUMN,
            PASSIVE_GRADE_FIRST_COLUMN,
            attr,
            passive_name,
            owned_grades,
            stats,
        )
        worksheet.cell(row_number, ACTIVE_NAME_COLUMN).alignment = LEFT_ALIGNMENT
        worksheet.cell(row_number, PASSIVE_NAME_COLUMN).alignment = LEFT_ALIGNMENT
        for column_number in range(1, 16):
            if column_number == 1:
                continue
            worksheet.cell(row_number, column_number).alignment = CENTER_ALIGNMENT
        row_number += 1

    worksheet.merge_cells(
        start_row=section_start_row,
        start_column=1,
        end_row=row_number - 1,
        end_column=1,
    )
    series_cell = worksheet.cell(section_start_row, 1)
    series_cell.value = series_name
    series_cell.fill = SERIES_FILL
    series_cell.font = Font(bold=True)
    series_cell.alignment = CENTER_ALIGNMENT
    return row_number


def _write_extra_skills_section(
    worksheet,
    start_row: int,
    owned_grades: dict[str, set[str]],
) -> int:
    known_keys = _known_series_keys()
    extra_base_names = sorted(
        base_name
        for base_name in owned_grades
        if base_name not in known_keys
    )
    if not extra_base_names:
        return start_row

    worksheet.cell(start_row, 1).value = "其他仙术"
    worksheet.cell(start_row, 2).value = "等级"
    worksheet.cell(start_row, 9).value = "已习得"
    for column_number in range(1, 16):
        worksheet.cell(start_row, column_number).fill = HEADER_FILL
        worksheet.cell(start_row, column_number).alignment = CENTER_ALIGNMENT
    worksheet.merge_cells(
        start_row=start_row,
        start_column=2,
        end_row=start_row,
        end_column=8,
    )
    worksheet.merge_cells(
        start_row=start_row,
        start_column=9,
        end_row=start_row,
        end_column=15,
    )

    row_number = start_row + 1
    for base_name in extra_base_names:
        grades = sorted(owned_grades[base_name])
        worksheet.cell(row_number, 2).value = "·".join(grades)
        worksheet.cell(row_number, 3).value = base_name
        worksheet.cell(row_number, 9).value = "✔"
        worksheet.cell(row_number, 2).alignment = CENTER_ALIGNMENT
        worksheet.cell(row_number, 3).alignment = LEFT_ALIGNMENT
        worksheet.cell(row_number, 9).alignment = CENTER_ALIGNMENT
        for column_number in range(4, 8):
            worksheet.cell(row_number, column_number).fill = EMPTY_FILL
        row_number += 1
    return row_number


def skills_to_excel(
    skills_path: str | Path = DEFAULT_SKILLS_PATH,
    template_path: str | Path | None = None,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> Path:
    """从 skills.txt 直接生成仙术持有表，不依赖任何 Excel 模板。

    template_path 仅为兼容旧调用保留并忽略；行内所有系列/属性/技能名
    均由内置 SERIES_META 骨架与 skills.txt 数据生成。skills.txt 中存在而
    骨架未收录的技能自动汇入「其他仙术」区，保证数据不丢失。
    """
    resolved_skills_path = Path(skills_path).resolve()
    resolved_output_path = Path(output_path).resolve()

    if not resolved_skills_path.is_file():
        raise FileNotFoundError(f"找不到仙术记录文件: {resolved_skills_path}")

    owned_grades = load_owned_grades(resolved_skills_path)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    _style_header(worksheet)
    stats = [0, 0]  # 持有标记数 / 浅蓝连续持有数
    row_number = 2
    for series_name, series_rows in SERIES_META:
        row_number = _write_series_section(
            worksheet,
            row_number,
            series_name,
            series_rows,
            owned_grades,
            stats,
        )
        row_number += 1
    row_number = _write_extra_skills_section(worksheet, row_number, owned_grades)

    resolved_output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_output_path = resolved_output_path.with_name(
        f"{resolved_output_path.stem}.tmp{resolved_output_path.suffix}"
    )
    try:
        workbook.save(temporary_output_path)
        os.replace(temporary_output_path, resolved_output_path)
    finally:
        if temporary_output_path.exists():
            temporary_output_path.unlink()

    print(
        f"仙术持有表已生成: {resolved_output_path}，"
        f"持有标记 {stats[0]} 项，浅蓝连续持有 {stats[1]} 项"
    )
    return resolved_output_path


if __name__ == "__main__":
    skills_to_excel()
